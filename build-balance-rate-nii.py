#!/usr/bin/env python3
"""
Three tables -- Balance, Rate, Net Interest Income -- with quarter-end dates as
columns and disclosed line items ("GL accounts") as rows.

    python3 build-balance-rate-nii.py

Input : nii-avgbal.csv (parse-nii-tables.py)
Output: balance-rate-nii-cof.xlsx  + three CSVs

ON "GL ACCOUNTS": SEC filings do not expose a general ledger. These are the
line items Capital One discloses in its average-balance table, which is the
finest account-level granularity available externally. They are groupings, not
GL account numbers.

BALANCE IS AN AVERAGE, NOT A PERIOD-END. The source table reports average
balances over the quarter. They will NOT tie to the balance sheet. The column
header is the quarter-end date purely to identify the period.

Q4 IS DERIVED, and legitimately so. Two identities were tested on the actual
data and both pass 48/48:

    interest is additive          6M interest = Q1 + Q2
    balances weight by months     6M avg = (Q1 avg x3 + Q2 avg x3) / 6

So Q4 follows:
    Q4 interest = FY interest - 9M interest
    Q4 balance  = (FY avg x 12 - 9M avg x 9) / 3
    Q4 rate     = Q4 interest / Q4 balance x 4      (annualised)

Rates are NOT derived by subtraction -- they are recomputed from the derived
interest and balance, because rates do not subtract.
"""
import csv, os, sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ARIAL = "Arial"
BLUE = Font(name=ARIAL, size=10, color="0000FF")
GREEN = Font(name=ARIAL, size=10, color="008000")
BLACK = Font(name=ARIAL, size=10)
BOLD = Font(name=ARIAL, size=10, bold=True)
BOLD_U = Font(name=ARIAL, size=10, bold=True, underline="single")
TITLE = Font(name=ARIAL, size=12, bold=True)
NOTE = Font(name=ARIAL, size=9, italic=True, color="595959")
HDR = PatternFill("solid", fgColor="D9E1F2")
CHK = PatternFill("solid", fgColor="FFF2CC")
TOP = Border(top=Side(style="thin", color="BFBFBF"))

MONEY = '$#,##0;($#,##0);-'
PCT = '0.00%;(0.00%);-'
CHKF = '#,##0.0;(#,##0.0);-'

# Presentation order. Label in the filing -> row label here.
ASSETS = [
    ("Credit card", "Credit card"),
    ("Consumer banking", "Consumer banking"),
    ("Commercial banking", "Commercial banking"),
    ("Total loans, including loans held for sale", "Total loans (incl. HFS)"),
    ("Investment securities", "Investment securities"),
    ("Cash equivalents and other interest-earning assets",
     "Cash equivalents and other"),
]
ASSET_TOTAL = ("Total interest-earning assets", "Total interest-earning assets")
LIABS = [
    ("Interest-bearing deposits", "Interest-bearing deposits"),
    ("Securitized debt obligations", "Securitized debt obligations"),
    ("Senior and subordinated notes", "Senior and subordinated notes"),
    ("Other borrowings and interest-bearing liabilities", "Other borrowings"),
]
LIAB_TOTAL = ("Total interest-bearing liabilities",
              "Total interest-bearing liabilities")

QEND = {1: "03-31", 2: "06-30", 3: "09-30", 4: "12-31"}


def qkey(p):
    return int(p[:4]) * 4 + int(p[-1])


def qdate(p):
    return f"{p[:4]}-{QEND[int(p[-1])]}"


def load():
    v = defaultdict(dict)          # (line_item, field) -> (period, months) -> val
    for r in csv.DictReader(open("nii-avgbal.csv")):
        v[(r["line_item"], r["field"])][(r["period"], int(r["months"]))] = \
            float(r["value"])
    return v


def derive_q4(v):
    """Fill Q4 from FY and 9M. Returns set of derived period labels."""
    derived = set()
    items = {li for li, f in v}
    for li in items:
        for yr in ("2023", "2024", "2025"):
            fy_i = v.get((li, "interest"), {}).get((f"{yr}FY", 12))
            n_i = v.get((li, "interest"), {}).get((f"{yr}Q3", 9))
            fy_b = v.get((li, "avg_balance"), {}).get((f"{yr}FY", 12))
            n_b = v.get((li, "avg_balance"), {}).get((f"{yr}Q3", 9))
            if None in (fy_i, n_i, fy_b, n_b):
                continue
            q4_i = fy_i - n_i
            q4_b = (fy_b * 12 - n_b * 9) / 3
            v[(li, "interest")][(f"{yr}Q4", 3)] = round(q4_i, 1)
            v[(li, "avg_balance")][(f"{yr}Q4", 3)] = round(q4_b, 1)
            if q4_b:
                # recompute the rate; rates cannot be subtracted
                v[(li, "yield_pct")][(f"{yr}Q4", 3)] = round(
                    q4_i / q4_b * 4 * 100, 2)
            derived.add(f"{yr}Q4")
    return derived


def sheet(wb, name, title, subtitle, periods, derived, fmt):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = TITLE
    ws["A2"] = subtitle; ws["A2"].font = NOTE
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 28
    ws.cell(row=4, column=1, value="Line item (as disclosed)").font = BOLD
    for i, p in enumerate(periods):
        c = ws.cell(row=4, column=2 + i,
                    value=qdate(p) + (" (D)" if p in derived else ""))
        c.font, c.fill = BOLD, HDR
        c.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 38
    for i in range(len(periods)):
        ws.column_dimensions[get_column_letter(2 + i)].width = 13
    ws.freeze_panes = ws.cell(row=5, column=2)
    return ws


def emit_block(ws, r, heading, rows, v, field, periods, fmt, scale=1.0):
    ws.cell(row=r, column=1, value=heading).font = BOLD_U
    r += 1
    first = r
    rowmap = {}
    for src, lbl in rows:
        if (src, field) not in v:
            continue
        ws.cell(row=r, column=1, value="  " + lbl).font = BLACK
        for i, p in enumerate(periods):
            val = v[(src, field)].get((p, 3))
            if val is not None:
                val = val * scale
            c = ws.cell(row=r, column=2 + i, value=val)
            c.font, c.number_format = BLUE, fmt
        rowmap[src] = r
        r += 1
    return r, first, rowmap


def build(wb, v, periods, derived, field, name, title, subtitle, fmt,
          scale=1.0):
    ws = sheet(wb, name, title, subtitle, periods, derived, fmt)
    r = 5
    r, a_first, a_map = emit_block(ws, r, "INTEREST-EARNING ASSETS",
                                   ASSETS, v, field, periods, fmt, scale)
    # reported total as its own row
    ws.cell(row=r, column=1, value="  " + ASSET_TOTAL[1]).font = BOLD
    for i, p in enumerate(periods):
        val = v.get((ASSET_TOTAL[0], field), {}).get((p, 3))
        if val is not None:
            val *= scale
        c = ws.cell(row=r, column=2 + i, value=val)
        c.font, c.number_format, c.border = BOLD, fmt, TOP
    a_total_row = r
    r += 2

    r, l_first, l_map = emit_block(ws, r, "INTEREST-BEARING LIABILITIES",
                                   LIABS, v, field, periods, fmt, scale)
    ws.cell(row=r, column=1, value="  " + LIAB_TOTAL[1]).font = BOLD
    for i, p in enumerate(periods):
        val = v.get((LIAB_TOTAL[0], field), {}).get((p, 3))
        if val is not None:
            val *= scale
        c = ws.cell(row=r, column=2 + i, value=val)
        c.font, c.number_format, c.border = BOLD, fmt, TOP
    l_total_row = r
    r += 2
    return ws, r, a_total_row, l_total_row, a_map, l_map


def main():
    if not os.path.exists("nii-avgbal.csv"):
        sys.exit("missing nii-avgbal.csv -- run parse-nii-tables.py")
    v = load()
    derived = derive_q4(v)
    periods = sorted({p for d in v.values() for (p, m) in d if m == 3}, key=qkey)
    print(f"periods: {len(periods)}  {periods[0]}..{periods[-1]}")
    print(f"derived Q4: {sorted(derived)}")

    wb = Workbook(); wb.remove(wb.active)

    # ---------------- 1. BALANCE
    ws, r, atot, ltot, _, _ = build(
        wb, v, periods, derived, "avg_balance", "Balance",
        "Capital One — Average Balance by line item ($mm)",
        "AVERAGE balance over the quarter, not period-end — will not tie to the "
        "balance sheet. Column header is the quarter-end date. (D) = Q4 derived "
        "as (FY avg x12 − 9M avg x9)/3. Source: COF 10-Q/10-K average balance table.",
        MONEY)
    ws.cell(row=r, column=1, value="Net earning assets (assets − liabilities)").font = BOLD
    for i in range(len(periods)):
        col = get_column_letter(2 + i)
        c = ws.cell(row=r, column=2 + i, value=f"={col}{atot}-{col}{ltot}")
        c.font, c.number_format = BOLD, MONEY
    bal_asset_row, bal_liab_row = atot, ltot

    # ---------------- 2. RATE
    ws2, r2, atot2, ltot2, _, _ = build(
        wb, v, periods, derived, "yield_pct", "Rate",
        "Capital One — Average Yield / Rate by line item (annualised)",
        "Annualised yield on assets and cost of funds on liabilities, as "
        "disclosed. (D) = Q4 recomputed as derived interest / derived balance x4; "
        "rates cannot be derived by subtraction. Stored as fractions.",
        PCT, scale=0.01)
    ws2.cell(row=r2, column=1, value="Net interest spread (asset yield − funding cost)").font = BOLD
    for i in range(len(periods)):
        col = get_column_letter(2 + i)
        c = ws2.cell(row=r2, column=2 + i, value=f"={col}{atot2}-{col}{ltot2}")
        c.font, c.number_format = BOLD, PCT
    r2 += 1
    ws2.cell(row=r2, column=1, value="Net interest margin (NII / avg earning assets)").font = BOLD
    for i in range(len(periods)):
        col = get_column_letter(2 + i)
        c = ws2.cell(
            row=r2, column=2 + i,
            value=(f"=IFERROR(('Net Interest Income'!{col}{{nii}}*4)"
                   f"/Balance!{col}{bal_asset_row},\"\")"))
        c.font, c.number_format = BOLD, PCT
    nim_row = r2

    # ---------------- 3. NET INTEREST INCOME
    ws3, r3, atot3, ltot3, _, _ = build(
        wb, v, periods, derived, "interest", "Net Interest Income",
        "Capital One — Interest Income / Expense by line item ($mm, quarterly)",
        "Interest income on assets, interest expense on liabilities, as "
        "disclosed (expense shown positive). NII = total interest income − total "
        "interest expense. (D) = Q4 derived as FY − 9M; interest is additive so "
        "this is exact.",
        MONEY)
    ws3.cell(row=r3, column=1, value="NET INTEREST INCOME").font = BOLD
    for i in range(len(periods)):
        col = get_column_letter(2 + i)
        c = ws3.cell(row=r3, column=2 + i, value=f"={col}{atot3}-{col}{ltot3}")
        c.font, c.number_format, c.border = BOLD, MONEY, TOP
    nii_row = r3
    r3 += 2

    # cross-check against the independently extracted income statement
    ws3.cell(row=r3, column=1,
             value="CHECK: NII vs income statement (flows.csv)").font = BOLD_U
    r3 += 1
    isnii = {}
    if os.path.exists("flows.csv"):
        for rr in csv.DictReader(open("flows.csv")):
            if (rr["filer"] == "cof" and rr["concept"] == "net_interest_income"
                    and int(rr["months"]) == 3 and rr["quarter"]):
                isnii[rr["quarter"]] = float(rr["value_usd"]) / 1e6
    ws3.cell(row=r3, column=1, value="  Income statement NII (XBRL)").font = BLACK
    for i, p in enumerate(periods):
        c = ws3.cell(row=r3, column=2 + i, value=isnii.get(p))
        c.font, c.number_format = GREEN, MONEY
    is_row = r3
    r3 += 1
    ws3.cell(row=r3, column=1, value="  Difference").font = BOLD
    for i in range(len(periods)):
        col = get_column_letter(2 + i)
        c = ws3.cell(row=r3, column=2 + i,
                     value=f"=IF({col}{is_row}=\"\",\"\",{col}{nii_row}-{col}{is_row})")
        c.font, c.number_format, c.fill = BOLD, CHKF, CHK
    r3 += 2
    ws3.cell(row=r3, column=1,
             value="Blue = as disclosed. Green = from the XBRL income statement "
                   "(independent source). Black = formula. Difference should be "
                   "zero: the MD&A average-balance table and the income statement "
                   "must agree on NII.").font = NOTE

    # patch the NIM formula now that we know the NII row
    for i in range(len(periods)):
        col = get_column_letter(2 + i)
        ws2.cell(row=nim_row, column=2 + i).value = (
            f"=IFERROR(('Net Interest Income'!{col}{nii_row}*4)"
            f"/Balance!{col}{bal_asset_row},\"\")")

    # ---------------- notes
    wsn = wb.create_sheet("Notes & Sources")
    wsn.column_dimensions["A"].width = 108
    notes = [
        ("Balance / Rate / Net Interest Income — Capital One", TITLE),
        ("", NOTE),
        ('ON "GL ACCOUNTS"', BOLD_U),
        ("SEC filings do not expose a general ledger. These rows are the line "
         "items Capital One discloses in its average-balance table — the finest "
         "account-level granularity available externally. They are groupings, "
         "not GL account numbers.", BLACK),
        ("", NOTE),
        ("WHAT EACH TABLE HOLDS", BOLD_U),
        ("Balance — AVERAGE balance over the quarter, in $mm. Not period-end; "
         "it will not tie to the balance sheet.", BLACK),
        ("Rate — annualised yield (assets) and cost of funds (liabilities), "
         "stored as fractions and displayed as percentages.", BLACK),
        ("Net Interest Income — interest income on assets and interest expense "
         "on liabilities, with NII computed as the difference of the totals.", BLACK),
        ("", NOTE),
        ("SOURCE", BOLD_U),
        ("Capital One Financial Corp (CIK 0000927628), Forms 10-Q and 10-K, "
         "MD&A table 'Average Balances, Net Interest Income and Net Interest "
         "Margin'. These MD&A tables are NOT XBRL-tagged, so values are parsed "
         "from the rendered HTML.", BLACK),
        ("", NOTE),
        ("Q4 DERIVATION, AND WHY IT IS SOUND", BOLD_U),
        ("No 3-month Q4 average-balance table is published: the 10-K table is "
         "annual and the 10-Q tables are 3/6/9-month. Q4 columns are marked (D).", BLACK),
        ("Two identities were tested against the actual data, both passing 48 of "
         "48 cases: interest is additive (6M = Q1 + Q2), and average balances "
         "weight by months (6M avg = (Q1 avg x3 + Q2 avg x3) / 6).", BLACK),
        ("Therefore: Q4 interest = FY − 9M (exact). Q4 balance = (FY avg x12 − "
         "9M avg x9) / 3. Q4 rate is RECOMPUTED from the derived interest and "
         "balance, never subtracted — rates do not subtract.", BLACK),
        ("", NOTE),
        ("VALIDATION", BOLD_U),
        ("The Net Interest Income sheet cross-checks its computed NII against "
         "the income statement extracted independently from XBRL. The two come "
         "from different places in the filing by different methods, so agreement "
         "is meaningful.", BLACK),
        ("", NOTE),
        ("CONVENTIONS", BOLD_U),
        ("Blue = as disclosed in the filing. Green = independent XBRL source. "
         "Black = live formula. Yellow = check cell, should read zero.", BLACK),
    ]
    rr = 1
    for t, f in notes:
        c = wsn.cell(row=rr, column=1, value=t)
        c.font = f; c.alignment = Alignment(wrap_text=True, vertical="top")
        rr += 1

    out = "balance-rate-nii-cof.xlsx"
    wb.save(out)
    print(f"-> {out}")

    # ---------------- CSV companions
    def dump(field, path, scale=1.0):
        allrows = ASSETS + [ASSET_TOTAL] + LIABS + [LIAB_TOTAL]
        with open(path, "w", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["line_item"] + [qdate(p) for p in periods])
            for src, lbl in allrows:
                if (src, field) not in v:
                    continue
                row = [lbl]
                for p in periods:
                    x = v[(src, field)].get((p, 3))
                    row.append("" if x is None else round(x * scale, 4))
                w.writerow(row)
        print(f"-> {path}")

    dump("avg_balance", "table-balance-cof.csv")
    dump("yield_pct", "table-rate-cof.csv")
    dump("interest", "table-nii-cof.csv")

    # console check
    print("\nNII check (computed from avg-bal table vs XBRL income statement):")
    bad = 0
    for p in periods:
        a = v.get((ASSET_TOTAL[0], "interest"), {}).get((p, 3))
        l = v.get((LIAB_TOTAL[0], "interest"), {}).get((p, 3))
        if a is None or l is None or p not in isnii:
            continue
        d = (a - l) - isnii[p]
        flag = "ok" if abs(d) < 1.0 else "MISMATCH"
        if flag != "ok":
            bad += 1
        print(f"   {qdate(p)}  {flag:9s} table {a-l:>9,.0f}  "
              f"XBRL {isnii[p]:>9,.0f}  diff {d:>7,.1f}")
    print(f"   {'all periods tie' if not bad else str(bad) + ' MISMATCHES'}")


if __name__ == "__main__":
    main()
